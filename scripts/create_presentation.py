from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path('docs/EV_LOGISTICS_SERVICE_PRESENTATION.pptx')
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(14, 31, 53)
BLUE = RGBColor(22, 112, 192)
CYAN = RGBColor(45, 181, 210)
GREEN = RGBColor(35, 153, 112)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(244, 248, 252)
MID = RGBColor(96, 108, 122)
LINE = RGBColor(210, 221, 232)


def textbox(slide, x, y, w, h, text, size=18, color=NAVY, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = '맑은 고딕'
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return shape


def base_slide(title, number, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = LIGHT
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(.18))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    textbox(slide, .7, .52, 11.8, .55, title, 27, NAVY, True)
    if subtitle:
        textbox(slide, .72, 1.12, 11.6, .35, subtitle, 12, MID)
    textbox(slide, 12.2, 6.95, .55, .25, str(number), 10, MID, False, PP_ALIGN.RIGHT)
    return slide


def bullets(slide, items, x=.9, y=1.65, w=11.5, h=4.9, size=19):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.name = '맑은 고딕'; p.font.size = Pt(size); p.font.color.rgb = NAVY
        p.space_after = Pt(15); p.level = 0
    return box


def card(slide, x, y, w, h, title, value, accent=BLUE, value_size=20):
    box = slide.shapes.add_shape(5, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid(); box.fill.fore_color.rgb = WHITE; box.line.color.rgb = LINE
    stripe = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(.08), Inches(h))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = accent; stripe.line.fill.background()
    textbox(slide, x + .23, y + .18, w - .4, .3, title, 11, MID, True)
    textbox(slide, x + .23, y + .62, w - .4, h - .75, value, value_size, NAVY, True)


# 0. 표지
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid(); s.background.fill.fore_color.rgb = NAVY
textbox(s, .9, 1.65, 11.5, .9, '일일 물류 배송 조회 서비스', 40, WHITE, True)
textbox(s, .95, 2.75, 11.2, .55, '전기 물류 트럭 배차 · 구간별 에너지 예측 · 충전 계획 · ETA', 20, RGBColor(178, 218, 246))
textbox(s, .95, 5.85, 11, .4, '부산항 신선대 컨테이너터미널 → 한진인천컨테이너터미널(HJIT)', 14, WHITE)
textbox(s, .95, 6.35, 5, .3, 'Service Presentation', 11, RGBColor(140, 166, 191))

# 1. 문제와 서비스
s = base_slide('1. 서비스 개요', 1, '운행 전에 차량·배터리·충전·도착 정보를 한 번에 확인')
bullets(s, [
    '사원번호와 배차일만으로 당일 차량·화물·배터리 상태를 조회',
    '465.6km 고정 노선의 17개 구간을 각각 계산해 소비에너지와 SOC 예측',
    '안전 SOC 10%를 기준으로 충전 필요 구간과 예상 도착시간 제공',
    '운영 결과는 일일 배차기록 CSV·Excel로 저장',
])

# 2. 사용자 흐름
s = base_slide('2. 사용자 흐름', 2)
steps = [('1', '사원번호·배차일 조회'), ('2', '차량·화물·운행조건 확인'), ('3', '구간별 에너지·SOC 계산'), ('4', '충전·도착시간 확인'), ('5', '일일 배차기록 저장')]
for i, (n, label) in enumerate(steps):
    x = .55 + i * 2.55
    card(s, x, 2.05, 2.2, 2.05, f'STEP {n}', label, CYAN if i < 3 else GREEN, 16)
    if i < len(steps) - 1:
        textbox(s, x + 2.2, 2.75, .35, .4, '→', 20, BLUE, True, PP_ALIGN.CENTER)
textbox(s, .85, 5.15, 11.7, .5, '조회 중에는 사원번호와 배차일을 잠그고, 메인화면 복귀 시 입력 상태를 초기화합니다.', 15, MID, False, PP_ALIGN.CENTER)

# 3. 입력/자동 생성
s = base_slide('3. 입력값과 자동 배차 조건', 3)
card(s, .7, 1.55, 3.85, 1.45, '사용자 입력', '사원번호 8자리 · 배차일', BLUE, 17)
card(s, 4.75, 1.55, 3.85, 1.45, '운행조건', '차량 · 화물 박스 · 속도 · HVAC', CYAN, 17)
card(s, 8.8, 1.55, 3.85, 1.45, '자동 생성', 'SOC · 운전성향 · 날씨', GREEN, 17)
bullets(s, [
    '화물: 1박스 10kg, 최초 3~10박스 자동 배정',
    '속도: 기본 100km/h, 5km/h 단위 조정',
    '출발 SOC: 50~90% 중 5% 단위 자동 생성',
    '같은 사원번호와 배차일은 동일한 결과를 재현',
], x=.95, y=3.45, w=11.2, h=2.6, size=17)

# 4. 핵심 화면
s = base_slide('4. 조회 결과 화면', 4, '운영자가 먼저 확인해야 할 정보를 한 화면에 배치')
cards = [
    ('사원번호', '12345678'), ('차량번호', '7호차 (Volvo FH)'),
    ('출발 배터리', '65%'), ('화물', '8박스 / 80kg'),
    ('충전 휴게소', 'A ~ B 구간'), ('운행시간', '10:00 → 16:42'),
    ('예상 소비량', '537.1kWh'), ('예측 신뢰', 'R² 94.2% / ±0.89'),
]
for i, (t, v) in enumerate(cards):
    x = .65 + (i % 2) * 6.15
    y = 1.45 + (i // 2) * 1.35
    card(s, x, y, 5.75, 1.05, t, v, BLUE if i < 4 else GREEN, 16)

# 5. ML
s = base_slide('5. 머신러닝 모델', 5)
card(s, .75, 1.55, 3.75, 1.45, '모델', 'Linear Regression', BLUE, 19)
card(s, 4.8, 1.55, 3.75, 1.45, '설명력', 'R² 94.17%', GREEN, 22)
card(s, 8.85, 1.55, 3.75, 1.45, '예상 오차', 'RMSE ±0.887', CYAN, 22)
bullets(s, [
    '기본 변수: 속도, 화물량, 온도, HVAC, 경사, 운전성향, 타이어 공기압, 거리',
    '파생변수: 속도 제곱, 온도 편차, 적재량×경사, HVAC×온도 등',
    'StandardScaler 적용 후 저장 모델을 재사용해 서비스 시작시간 단축',
], x=.9, y=3.45, w=11.6, h=2.4, size=17)

# 6. 구간별 계산
s = base_slide('6. 구간별 거리·경사 에너지 계산', 6, '전체 평균으로 합치지 않고 17개 구간을 각각 예측')
bullets(s, [
    '휴게소 사이 구간마다 거리와 평균경사를 모델에 개별 입력',
    '구간별 kWh/100km → 구간 소비에너지 → SOC 순차 차감',
    '충전 가능 휴게소와 안전 SOC를 함께 검토해 충전 정차 판단',
], x=.8, y=1.45, w=6.1, h=2.7, size=17)
for i, (label, value) in enumerate([('계산 구간', '17개'), ('노선 거리', '465.6km'), ('경사 범위', '-0.279~0.600%'), ('테스트 소비량', '약 537.1kWh')]):
    x = 7.25 + (i % 2) * 2.75; y = 1.55 + (i // 2) * 1.65
    card(s, x, y, 2.45, 1.3, label, value, CYAN if i < 2 else GREEN, 17)
textbox(s, .85, 5.65, 11.7, .55, '경사도는 Copernicus DEM 지점 고도 차이로 산출한 구간 평균값이며, 순간 최대경사와는 다를 수 있습니다.', 13, MID, False, PP_ALIGN.CENTER)

# 7. SOC·충전·ETA
s = base_slide('7. SOC·충전·도착시간 계산', 7)
bullets(s, [
    '출발: 배차일 오전 10시, SOC 50~90% 자동 배정',
    '안전 기준: 구간 도착 후 최소 SOC 10% 유지',
    '충전 필요 시: 이용 가능한 휴게소에서 목표 SOC 100%',
    '충전시간: 충전량 ÷ 충전출력 ÷ 효율 92%',
    '총 소요시간: 주행시간 + 충전시간 + 진입·연결·출차 7분',
], size=18)

# 8. 데이터 출처/한계
s = base_slide('8. 노선 경사 데이터 조사', 8)
bullets(s, [
    '기존 노선 CSV에는 거리·충전기 정보만 있고 경사·고도 데이터가 없었음',
    '위치: OpenStreetMap Nominatim / 고도: Copernicus DEM via Open-Meteo',
    '서울방향 휴게소 좌표를 재검증하고 출발·도착 주소 기준으로 보정',
    '공식 도로 종단선형을 찾지 못해 지점 고도 차이 기반 평균경사 적용',
    '향후 실제 GPS·고도 로그 또는 도로 중심선 종단 프로파일로 교체 가능',
], size=17)

# 9. 운영 기능
s = base_slide('9. 운영 화면과 기록 관리', 9)
bullets(s, [
    '핵심 결과 아래에 배차 상세정보 · 에너지·충전 정보 · 노선 정보 · 충전소 정보 제공',
    '일일 배차기록: 사원번호 · 차량번호 · 배차 허가 여부만 저장',
    '파일명에 배차일을 포함하고 CSV·Excel 두 형식 지원',
    '사원번호 자리수·숫자 검증, 차량 중복·예비차량·덮어쓰기 처리',
], size=18)

# 10. 발표용 화면 자리
s = base_slide('10. 실제 사용화면', 10, '발표 직전 Streamlit 화면 캡처를 아래 영역에 교체 삽입')
ph = s.shapes.add_shape(5, Inches(.8), Inches(1.45), Inches(11.75), Inches(4.95))
ph.fill.solid(); ph.fill.fore_color.rgb = WHITE; ph.line.color.rgb = BLUE; ph.line.width = Pt(2)
textbox(s, 2.0, 3.35, 9.3, .65, '[ Streamlit 운행 정보 화면 스크린샷 ]', 23, MID, True, PP_ALIGN.CENTER)
textbox(s, 2.0, 4.15, 9.3, .4, '권장 캡처: 운행 정보 + 충전 추천 + 운행시간 + 예측 신뢰 지표', 12, MID, False, PP_ALIGN.CENTER)

# 11. 결론
s = base_slide('11. 기대 효과 및 확장 방향', 11)
bullets(s, [
    '배차 전에 차량 조건과 배터리 도착 가능성을 빠르게 확인',
    '구간별 경사와 거리 기반 예측으로 위험 구간과 충전 필요성 파악',
    '반복 학습 없이 저장 모델을 사용해 빠른 서비스 응답 제공',
    '향후 실차 운행 로그·실시간 교통·충전소 혼잡 데이터로 고도화',
], size=19)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'created: {OUT} / slides: {len(prs.slides)}')
